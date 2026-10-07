from io import BytesIO


def test_evidence_brief_preserves_unknown_recommendation_and_source_agent():
    from app.core.reports.evidence_brief import build_evidence_brief

    brief = build_evidence_brief(
        {"id": "deal-1", "final_score": None},
        [{
            "agent_type": "financial_analyst",
            "data": {
                "key_findings": ["Reported ARR is $12m"],
                "data_gaps": ["No audited historical revenue supplied"],
                "citations": [{"title": "Management accounts", "page": 4}],
            },
        }],
    )

    assert brief["recommendation"] is None
    assert brief["score_recorded"] is False
    assert brief["successful_analysis_count"] == 1
    assert brief["agents_with_citations"] == 1
    assert brief["findings"][0]["agent"] == "financial_analyst"
    assert brief["unknowns"][0]["text"] == "No audited historical revenue supplied"


def test_evidence_brief_includes_nested_financial_data_limitations():
    from app.core.reports.evidence_brief import build_evidence_brief

    brief = build_evidence_brief(
        {"id": "deal-2"},
        [{
            "agent_type": "financial_analyst",
            "data": {"financial_analysis": {"data_limitations": ["Cash flow not supplied", "No valuation estimated"]}},
        }],
    )

    assert [item["text"] for item in brief["unknowns"]] == ["Cash flow not supplied", "No valuation estimated"]


def test_evidence_brief_curates_periodic_api_metrics_with_stable_source_ids():
    from app.core.reports.evidence_brief import build_evidence_brief

    brief = build_evidence_brief(
        {"id": "deal-3"},
        [{
            "agent_type": "financial_analyst", "data": {
                "sources": [{"title": "Issuer 10-K", "period": "2025", "url": "https://example.test/10k"}],
                "historical_financials": [{
                    "period": "2025", "fiscal_year": "FY2025", "period_end_date": "2025-12-31",
                    "revenue": 2_500_000_000, "net_income": None,
                    "revenue_yoy_percent": 12.5, "operating_margin_percent": 4.2,
                    "source_url": "https://example.test/10k",
                }],
            },
        }],
    )

    revenue = next(item for item in brief["data_points"] if item["metric"] == "revenue")
    assert revenue["value"] == 2_500_000_000
    assert revenue["period"] == "FY2025"
    assert revenue["source_id"] == "S1"
    assert revenue["basis"] == "reported_source_linked"
    growth = next(item for item in brief["data_points"] if item["metric"] == "revenue_yoy_percent")
    assert growth["basis"] == "derived_source_linked"
    margin = next(item for item in brief["data_points"] if item["metric"] == "operating_margin_percent")
    assert margin["basis"] == "derived_source_linked"
    assert brief["sources"][0]["url"] == "https://example.test/10k"
    assert not any(item["metric"] == "net_income" for item in brief["data_points"])


def test_document_synthesis_receives_curated_evidence_not_raw_tool_payload():
    import asyncio
    from types import SimpleNamespace
    from app.core.reports.document_planner import prepare_document_payload

    class FakeAgent:
        async def run(self, task, context):
            pack = context["deal_data"]
            assert pack["data_points"][0]["source_id"] == "S1"
            assert "raw_data" not in str(pack)
            return SimpleNamespace(success=True, data={"key_takeaways": []})

    class FakeRegistry:
        def get(self, name):
            return FakeAgent() if name == "business_analyst" else None

    activities = [{
        "agent_type": "financial_analyst", "success": True,
        "data": {
            "sources": [{"title": "Filing", "url": "https://example.test/filing"}],
            "historical_financials": [{
                "period": "2025", "fiscal_year": "FY2025", "revenue": 100,
                "source_url": "https://example.test/filing",
            }],
            "raw_data": {"large": "unneeded tool payload"},
        },
    }]
    _, analyst_data, evidence = asyncio.run(prepare_document_payload(FakeRegistry(), {"id": "deal"}, activities))
    assert evidence["data_points"][0]["value"] == 100
    assert analyst_data["_document_qa"]["status"] == "ready_for_review"


def test_synthesis_curation_rejects_unsupported_recommendations_metrics_and_citations():
    from app.core.reports.document_planner import _curate_synthesis

    evidence = {
        "data_points": [{"agent": "financial_analyst", "metric": "revenue", "value": 100,
                         "period": "FY2025", "source_id": "S1"}],
        "sources": [{"id": "S1", "url": "https://example.test/source"}],
        "findings": [], "unknowns": [], "recommendation": None,
    }
    output = _curate_synthesis({
        "executive_summary": {"answer": "Proceed with the acquisition."},
        "key_takeaways": [{"title": "Growth", "description": "Revenue doubled [S99]."}],
        "financial_synthesis": {"narrative": "Revenue reached $999bn [S1].", "key_metrics": {"EBITDA": 500}},
        "risk_matrix": [{"severity": "high"}],
        "action_items": ["Proceed immediately"],
    }, evidence)

    assert output["executive_summary"]["answer"] == "No recommendation recorded; human review required."
    assert output["key_takeaways"] == []
    assert output["financial_synthesis"]["key_metrics"] == {"Revenue FY2025": "100 [S1]"}
    assert output["risk_matrix"] == []
    assert output["action_items"] == []
    assert "$999bn" not in output["financial_synthesis"]["narrative"]


def test_deterministic_financial_narrative_derives_perioded_trends_and_cites_them():
    from app.core.reports.document_planner import _deterministic_synthesis

    evidence = {
        "data_points": [
            {"agent": "financial_analyst", "metric": "revenue", "value": 100, "period": "FY2023", "source_id": "S1"},
            {"agent": "financial_analyst", "metric": "revenue", "value": 121, "period": "FY2024", "source_id": "S2"},
            {"agent": "financial_analyst", "metric": "revenue", "value": 160, "period": "FY2025", "source_id": "S3"},
            {"agent": "financial_analyst", "metric": "revenue_yoy_percent", "value": 21, "period": "FY2024", "source_id": "S2"},
            {"agent": "financial_analyst", "metric": "revenue_yoy_percent", "value": 32.23, "period": "FY2025", "source_id": "S3"},
            {"agent": "financial_analyst", "metric": "operating_margin_percent", "value": 20, "period": "FY2023", "source_id": "S1"},
            {"agent": "financial_analyst", "metric": "operating_margin_percent", "value": 25, "period": "FY2025", "source_id": "S3"},
        ],
        "findings": [], "unknowns": [], "recommendation": None,
    }
    result = _deterministic_synthesis(evidence)
    narrative = result["financial_synthesis"]["narrative"]
    takeaways = " ".join(item["description"] for item in result["key_takeaways"])

    assert "CAGR" in narrative
    assert "FY2023" in narrative and "FY2025" in narrative
    assert "accelerated" in takeaways
    assert "Operating margin" in " ".join(item["title"] for item in result["key_takeaways"])
    assert "[S1]" in takeaways and "[S3]" in takeaways


def test_empty_financial_inputs_do_not_create_model_assumptions_or_recommendation():
    from openpyxl import load_workbook
    from app.core.reports.report_generator import generate_excel

    content = generate_excel(
        {"target_company": "TestCo", "industry": "software", "final_score": 0.9},
        {"financial_synthesis": {"key_metrics": {"Revenue ($M)": 100}}},
        [],
    )
    workbook = load_workbook(BytesIO(content), data_only=False)

    assert "No recommendation recorded" in workbook["Executive Summary"]["B6"].value
    assert workbook["Income Statement"]["A2"].value == "No structured financial data points recorded"
    assert workbook["DCF Analysis"]["B3"].value == "Not calculated"
    assert workbook["LBO Returns"]["B3"].value == "Not calculated"
    assert "No sourced comparable" in workbook["Comps"]["A2"].value
    assert "No structured risk" in workbook["Risk Matrix"]["A2"].value
    assert not any(
        isinstance(cell.value, str) and cell.value.startswith("=")
        for row in workbook["Income Statement"].iter_rows()
        for cell in row
    )


def test_docx_does_not_render_llm_metrics_as_reported_or_synthetic_forecast():
    from docx import Document
    from app.core.reports.report_generator import generate_docx

    content = generate_docx(
        {"target_company": "TestCo", "industry": "software"},
        {"financial_synthesis": {"key_metrics": {"Revenue ($M)": 100}}},
        [],
    )
    doc = Document(BytesIO(content))
    text = "\n".join(p.text for p in doc.paragraphs)
    text += "\n" + "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)

    assert "No structured financial metrics were recorded" not in text
    assert "Forecast status" not in text
    assert "5-Year Projection Summary" not in text
    assert "100.0" not in text


def test_docx_renders_curated_api_points_and_source_register():
    from docx import Document
    from app.core.reports.report_generator import generate_docx

    content = generate_docx(
        {"id": "deal-4", "name": "Example review", "target_company": "Example Co"},
        {
            "executive_summary": {"situation": "FY2025 filing reports revenue."},
            "key_takeaways": [{"title": "Revenue", "description": "Revenue was $2.5bn [S1]."}],
            "financial_synthesis": {"narrative": "Revenue was reported for FY2025 [S1]."},
            "_evidence_brief": {
                "successful_analysis_count": 1, "agents_with_citations": 1,
                "findings": [], "unknowns": [], "notice": "Recorded outputs only.",
                "data_points": [{"metric": "revenue", "value": 2_500_000_000,
                                 "period": "FY2025", "basis": "reported", "source_id": "S1"}],
                "sources": [{"id": "S1", "title": "Example 10-K", "period": "FY2025",
                             "url": "https://example.test/10k", "filed": "2026-02-01"}],
            },
        }, [],
    )
    report = Document(BytesIO(content))
    text = "\n".join(p.text for p in report.paragraphs)
    text += "\n" + "\n".join(cell.text for table in report.tables for row in table.rows for cell in row.cells)
    assert "Decision-Relevant Takeaways" in text
    assert "Table of Contents" in text
    assert "Financial Metrics and Source Status" in text
    assert "Update this field in Word" not in text
    assert "Evidence basis" in text
    assert "$2.5bn" in text
    assert "[S1] Example 10-K" in text
    assert "https://example.test/10k" in text


def test_docx_omits_unavailable_risk_and_visual_sections_and_uncurated_raw_metrics():
    from docx import Document
    from app.core.reports.report_generator import generate_docx

    content = generate_docx(
        {"target_company": "NoSource Co"},
        {"_evidence_brief": {"successful_analysis_count": 1, "sources": [], "data_points": [], "unknowns": []}},
        [{
            "agent_type": "financial_analyst",
            "success": True,
            "data": {"revenue": 900_000_000, "profitability": {"gross_margin": 0.68}},
            "summary": "Values are uncited agent output.",
        }],
    )
    doc = Document(BytesIO(content))
    text = "\n".join(p.text for p in doc.paragraphs)
    text += "\n" + "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)

    assert "Risk Matrix" not in text and "Risk Assessment" not in text
    assert "Visual Analytics Annex" not in text and "Visual Analysis" not in text
    assert "900.0" not in text and "$900.0m" not in text


def test_cash_flow_points_without_citations_are_not_labeled_reported():
    from app.core.reports.evidence_brief import build_evidence_brief

    brief = build_evidence_brief({"id": "deal-cash"}, [{
        "agent_type": "financial_analyst", "success": True,
        "data": {"cash_flow": {"operating_cash_flow": 100, "free_cash_flow": 80}},
    }])

    assert {point["metric"]: point["basis"] for point in brief["data_points"]} == {
        "operating_cash_flow": "unverified_no_source",
        "free_cash_flow": "derived_unverified_no_source",
    }
    assert len(brief["unknowns"]) == 2
    assert all("no linked source citation" in gap["text"] for gap in brief["unknowns"])
    assert all("fiscal period not established" in gap["text"] for gap in brief["unknowns"])


def test_report_docx_includes_chat_context_without_treating_it_as_evidence():
    from docx import Document
    from app.core.reports.report_generator import generate_docx

    content = generate_docx(
        {"id": "deal-chat", "target_company": "Example Co"},
        {
            "_chat_context": {
                "mandate": "Assess revenue quality and customer concentration.",
                "chat_summary": "Revenue is $900m and concentration is high.",
            },
            "_document_qa": {"status": "review_required"},
            "_evidence_brief": {"successful_analysis_count": 0, "sources": [], "data_points": [], "unknowns": []},
        },
        [],
    )
    doc = Document(BytesIO(content))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Assess revenue quality and customer concentration." in text
    assert "Revenue is $900m" in text
    assert "not independently verified" in text
    assert "source records: 0" in text
    assert "Document Plan & Quality Notes" not in text


def test_pptx_title_bars_use_short_target_subtitles():
    from io import BytesIO
    from pptx import Presentation
    from app.core.reports.report_generator import generate_pptx

    content = generate_pptx(
        {"id": "deal-title", "name": "A very long consulting engagement name that must not overflow the report slide title bar", "target_company": "Example Co"},
        {"_evidence_brief": {"sources": [{
            "id": "S1", "period": "FY2026", "filed": "2026-07-29", "form": "10-K",
            "title": "Example 10-K", "url": "https://www.sec.gov/Archives/edgar/data/1/filing/",
        }]}},
        [],
    )
    presentation = Presentation(BytesIO(content))
    title_bar_text = [
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.has_text_frame and shape.text.startswith("Executive Summary")
    ]
    assert title_bar_text == ["Executive Summary\nTarget: Example Co"]
    source_text = "\n".join(
        shape.text for slide in presentation.slides for shape in slide.shapes if shape.has_text_frame
    )
    assert "[S1] Example 10-K" in source_text
    assert "https://www.sec.gov/" in source_text


def test_excel_financial_schedule_uses_curated_points_and_displays_source_status():
    from openpyxl import load_workbook
    from app.core.reports.report_generator import generate_excel

    brief = {
        "data_points": [{
            "metric": "operating_cash_flow", "value": 182935000000,
            "period": "Not established", "basis": "unverified_no_source", "source_id": None,
        }, {
            "metric": "operating_margin_percent", "value": 46.78,
            "period": "FY2026", "basis": "reported_source_linked", "source_id": "S1",
        }],
        "sources": [{
            "id": "S1", "period": "FY2026", "filed": "2026-07-29", "form": "10-K",
            "title": "Example SEC filing", "url": "https://www.sec.gov/Archives/edgar/data/1/filing/",
        }],
    }
    content = generate_excel(
        {"target_company": "Example Co"},
        {"_evidence_brief": brief},
        [{"agent_type": "financial_analyst", "data": {"revenue": 331839000000}}],
    )
    sheet = load_workbook(BytesIO(content), data_only=False)["Income Statement"]

    assert [sheet.cell(1, col).value for col in range(1, 6)] == [
        "Metric", "Value (USD unless %)", "Period", "Evidence basis", "Source ID",
    ]
    assert sheet["B2"].value == 46.78
    assert sheet["B2"].number_format == '0.00"%"'
    assert sheet["D2"].value == "reported_source_linked"
    assert sheet["E2"].value == "S1"
    assert sheet["B3"].value == 182935000000
    assert sheet["B3"].number_format == '$#,##0.0,,,"bn";($#,##0.0,,,"bn");-'
    assert sheet["D3"].value == "unverified_no_source"
    assert sheet["E3"].value == "No citation recorded"
    sources = load_workbook(BytesIO(content), data_only=False)["Sources & References"]
    assert sources["A2"].value == "S1"
    assert sources["B2"].value == "FY2026"
    assert sources["F2"].hyperlink.target == "https://www.sec.gov/Archives/edgar/data/1/filing/"


def test_pdf_includes_latest_financial_metrics_and_clickable_source_reference():
    from pypdf import PdfReader
    from app.core.reports.report_generator import generate_pdf

    content = generate_pdf(
        {"id": "deal-pdf-evidence", "target_company": "Example Co"},
        {"_evidence_brief": {
            "successful_analysis_count": 1,
            "agents_with_citations": 1,
            "notice": "Recorded outputs only.",
            "data_points": [
                {"metric": "revenue", "value": 100_000_000_000, "period": "FY2026", "basis": "reported_source_linked", "source_id": "S1"},
                {"metric": "free_cash_flow", "value": 10_000_000_000, "period": "FY2026", "basis": "derived_source_linked", "source_id": "S1"},
                {"metric": "revenue", "value": 80_000_000_000, "period": "FY2025", "basis": "reported_source_linked", "source_id": "S2"},
            ],
            "sources": [{"id": "S1", "title": "Example 10-K", "url": "https://www.sec.gov/Archives/edgar/data/1/filing/"}],
            "findings": [], "unknowns": [],
        }},
        [],
    )
    pdf = PdfReader(BytesIO(content))
    text = "\n".join(page.extract_text() or "" for page in pdf.pages)

    assert "Financial Metrics & Evidence" in text
    assert "Revenue" in text and "$100.0bn" in text
    assert "Free Cash Flow" in text and "FY2026" in text
    assert "Example 10-K" in text and "S1" in text
    links = [
        annotation.get_object().get("/A", {}).get("/URI")
        for page in pdf.pages
        for annotation in page.get("/Annots", [])
        if annotation.get_object().get("/A")
    ]
    assert any(str(link).startswith("https://www.sec.gov/") for link in links)


def test_chat_context_is_selected_by_deal_and_extracts_latest_synthesis():
    from app.core.reports.document_planner import extract_chat_context

    conversations = [
        {"id": "other", "dealId": "other-deal", "updatedAt": 10, "messages": [
            {"role": "user", "content": "Do not include this."},
        ]},
        {"id": "right", "dealId": "deal-1", "updatedAt": 20, "messages": [
            {"role": "user", "content": "Assess the target."},
            {"role": "agent", "agentName": "DealForge Summary", "content": "Summary one."},
            {"role": "agent", "metadata": {"_isSynthesis": True}, "content": "Latest saved summary."},
        ]},
    ]

    result = extract_chat_context(conversations, "deal-1")
    assert result == {"mandate": "Assess the target.", "chat_summary": "Latest saved summary."}


def test_report_inputs_use_all_persisted_tasks_and_ignore_unpersisted_activity():
    from types import SimpleNamespace
    from app.core.reports.document_planner import build_report_agent_results

    task_lists = [SimpleNamespace(items=[
        SimpleNamespace(
            id="task-1", assigned_agent="financial_analyst", title="Financials",
            status="done", updated_at="2026-10-02T10:00:00Z", result={"revenue": 100},
        ),
        SimpleNamespace(
            id="task-2", assigned_agent="risk_assessor", title="Risks",
            status="done", updated_at="2026-10-02T10:01:00Z", result={"risks": ["Risk A"]},
        ),
    ])]
    activities = [
        {"agent_type": "financial_analyst", "data": {"revenue": 100}, "provider": "lmstudio"},
        {"agent_type": "market_researcher", "data": {"finding": "Unpersisted old case"}},
    ]

    results = build_report_agent_results(task_lists, activities)

    assert [(item["agent_type"], item["data"]) for item in results] == [
        ("financial_analyst", {"revenue": 100}),
        ("risk_assessor", {"risks": ["Risk A"]}),
    ]
    assert results[0]["provider"] == "lmstudio"


def test_report_requires_human_review_when_metrics_have_no_source_records():
    import asyncio
    from app.core.reports.document_planner import prepare_document_payload

    _, report_data, evidence = asyncio.run(prepare_document_payload(None, {"id": "deal-uncited"}, [{
        "agent_type": "financial_analyst",
        "success": True,
        "data": {"historical_financials": [{"fiscal_year": "FY2025", "revenue": 12_000_000}]},
    }]))

    assert evidence["sources"] == []
    assert evidence["data_points"][0]["basis"] == "unverified_no_source"
    assert report_data["_document_qa"]["status"] == "review_required"
    assert any("no linked source citation" in warning for warning in report_data["_document_qa"]["warnings"])

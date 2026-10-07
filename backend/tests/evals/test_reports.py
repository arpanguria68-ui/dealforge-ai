"""Validate generated deliverable artifacts, not synthetic expected prose."""
from io import BytesIO

from app.core.reports.report_generator import generate_docx, generate_excel, generate_pdf, generate_pptx


def _sample():
    return (
        {"name": "Eval Target", "target_company": "Eval Target", "tenant_id": "default"},
        {"executive_summary": {"situation": "Review of supplied, fictional case inputs."}},
        [],
    )


def test_docx_is_valid_and_contains_deal_identity():
    from docx import Document

    deal, analysis, results = _sample()
    payload = generate_docx(deal, analysis, results)
    document = Document(BytesIO(payload))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert payload.startswith(b"PK")
    assert "Eval Target" in text


def test_excel_is_valid_workbook_with_expected_tabs():
    from openpyxl import load_workbook

    deal, analysis, results = _sample()
    payload = generate_excel(deal, analysis, results)
    workbook = load_workbook(BytesIO(payload), read_only=True)
    assert "Executive Summary" in workbook.sheetnames
    assert len(workbook.sheetnames) >= 1


def test_pdf_and_pptx_are_valid_container_outputs():
    from pypdf import PdfReader
    from pptx import Presentation

    deal, analysis, results = _sample()
    pdf = generate_pdf(deal, analysis, results)
    pptx = generate_pptx(deal, analysis, results)
    assert len(PdfReader(BytesIO(pdf)).pages) >= 1
    presentation = Presentation(BytesIO(pptx))
    assert len(presentation.slides) >= 1

from io import BytesIO

from docx import Document

from app.core.reports.report_generator import generate_structured_analysis_docx


def test_structured_docx_preserves_results_and_marks_missing_values():
    content = generate_structured_analysis_docx(
        {"id": "deal-1", "target_company": "Acme Widgets", "final_score": None},
        {
            "status": "completed",
            "items": [
                {
                    "title": "Financial assessment",
                    "assigned_agent": "financial_analyst",
                    "status": "done",
                    "result": {
                        "revenue_analysis": {"arr": 10_000_000, "growth_rate": None},
                        "profitability": {"ebitda": 2_000_000},
                    },
                }
            ],
        },
    )

    report = Document(BytesIO(content))
    text = "\n".join(paragraph.text for paragraph in report.paragraphs)
    table_text = "\n".join(
        cell.text for table in report.tables for row in table.rows for cell in row.cells
    )

    assert "Acme Widgets" in text
    assert "not zero" in text
    assert "Not established" in table_text
    assert "10,000,000" in table_text
    assert "2,000,000" in table_text
    assert "18.5" not in table_text


def test_financial_docx_is_curated_and_omits_raw_sec_calculation_dump():
    content = generate_structured_analysis_docx(
        {"id": "deal-2", "target_company": "Microsoft Corporation"},
        {
            "status": "completed",
            "items": [{
                "title": "Financial assessment", "assigned_agent": "financial_analyst", "status": "done",
                "result": {
                    "data_source": "sec_edgar_companyfacts",
                    "confidence_basis": "not_calibrated_source_report",
                    "historical_financials": [{
                        "fiscal_year": "FY2026", "period_end_date": "2026-06-30",
                        "revenue": 331_839_000_000, "revenue_yoy_percent": 17.79,
                        "net_income": 133_749_000_000, "filing_form": "10-K",
                        "filing_date": "2026-07-29", "source_url": "https://sec.test/fy2026",
                    }],
                    "profitability": {"net_income": 133_749_000_000, "ebitda": None},
                    "cash_flow": {"operating_cash_flow": 182_935_000_000,
                                  "capital_expenditures": 115_948_000_000,
                                  "free_cash_flow": 66_987_000_000, "fiscal_year": "FY2026"},
                    "balance_sheet": {"cash": 20_935_000_000, "long_term_debt": 40_294_000_000},
                    "valuation": {"dcf_estimate": None, "multiple_estimate": None},
                    "sources": [{"period": "2026", "title": "Microsoft 10-K",
                                 "filed": "2026-07-29", "url": "https://sec.test/fy2026"}],
                    "calculations": [{"name": "fetch_financial_statements", "data": {"raw": "omit-me"}}],
                    "data_limitations": ["EBITDA not established."],
                },
            }],
        },
    )

    report = Document(BytesIO(content))
    text = "\n".join(paragraph.text for paragraph in report.paragraphs)
    table_text = "\n".join(
        cell.text for table in report.tables for row in table.rows for cell in row.cells
    )

    assert "Microsoft Corporation" in text
    assert "FY2026" in table_text
    assert "$331.8bn" in table_text
    assert "17.79%" in table_text
    assert "https://sec.test/fy2026" in text + table_text
    assert "$182.9bn" in table_text
    assert "omit-me" not in text + table_text
    assert "calculations" not in text + table_text


def test_financial_docx_includes_key_findings_and_compact_currency():
    content = generate_structured_analysis_docx(
        {"id": "deal-3", "target_company": "Example Co"},
        {"status": "completed", "items": [{
            "title": "Financial assessment", "assigned_agent": "financial_analyst", "status": "done",
            "result": {
                "historical_financials": [{
                    "period": "2025", "fiscal_year": "FY2025", "period_end_date": "2025-12-31",
                    "revenue": 2_500_000_000, "net_income": None, "source_url": "https://example.test/filing",
                }],
                "key_findings": [{"finding": "revenue_cagr", "text": "Revenue grew at 12.00% CAGR."}],
                "sources": [{"period": "2025", "title": "Example filing", "url": "https://example.test/filing"}],
                "data_limitations": [],
            },
        }]},
    )
    report = Document(BytesIO(content))
    text = "\n".join(paragraph.text for paragraph in report.paragraphs)
    table_text = "\n".join(cell.text for table in report.tables for row in table.rows for cell in row.cells)
    assert "Revenue grew at 12.00% CAGR." in text
    assert "$2.5bn" in table_text
    assert "Not established" in table_text

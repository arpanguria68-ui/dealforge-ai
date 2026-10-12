import io

from docx import Document

from app.core.reports.docx_rendering import add_markdown, format_value, render_payload
from app.core.reports.evidence_brief import build_evidence_brief


def _doc_text(doc):
    return [(p.style.name, p.text) for p in doc.paragraphs if p.text.strip()]


def test_markdown_becomes_real_word_structure():
    doc = Document()
    add_markdown(doc, "# Memo\n\nIntro with **bold** and *italic*.\n\n## Risks\n- first\n- second\n\n"
                      "| Metric | Value |\n|---|---|\n| Revenue | $457M |\n\n---\n1. step one\n")
    styles = dict((t, s) for s, t in _doc_text(doc))
    assert styles["Memo"].startswith("Heading") and styles["Risks"].startswith("Heading")
    assert styles["first"] == "List Bullet" and styles["step one"] == "List Number"
    assert len(doc.tables) == 1 and doc.tables[0].rows[1].cells[1].text == "$457M"
    assert not any(t.lstrip().startswith(("#", "**", "|")) for _, t in _doc_text(doc))
    intro = next(p for p in doc.paragraphs if p.text.startswith("Intro"))
    assert any(r.bold for r in intro.runs) and any(r.italic for r in intro.runs)


def test_chat_widget_markup_is_dropped():
    doc = Document()
    add_markdown(doc, 'Before\n<FinanceWidget id="x">\n```json\n{"a": 1}\n```\n</FinanceWidget>\nAfter')
    text = " ".join(t for _, t in _doc_text(doc))
    assert "FinanceWidget" not in text and "Before" in text and "After" in text


def test_value_formatting_uses_key_meaning():
    assert format_value("weighted_estimate", 934_449_473.89) == "$934.45M"
    assert format_value("wacc", 0.0958) == "9.58%"
    assert format_value("cost_of_equity", 0.1019) == "10.19%"
    assert format_value("gross_margin", 77.56) == "77.56%"
    assert format_value("ev_revenue_multiple", 2.594) == "2.59x"
    assert format_value("annual_revenue", 457_547_000) == "$457.55M"
    assert format_value("flag", True) == "Yes"


def test_payload_renders_findings_not_a_stub():
    doc = Document()
    render_payload(doc, {
        "revenue_analysis": {"annual_revenue": 457_547_000, "growth_rate": 12.72},
        "key_findings": ["Growth slowed to 12.7%", "Gross margin 77.6%"],
        "historical_financials": [
            {"fiscal_year": f"FY{y}", "revenue": 100e6 * (y - 2020), "gross_profit": 70e6, "operating_income": -5e6,
             "net_income": -6e6, "ebitda": 1e6, "period": str(y), "filing_form": "10-K", "filing_date": "2025-02-01"}
            for y in (2023, 2024, 2025)],
        "_validation": {"valid": True}, "provider_warning": "[Provider unavailable] x", "raw_response": "zzz",
    })
    text = " ".join(t for _, t in _doc_text(doc))
    assert "$457.55M" in text and "12.72%" in text and "Growth slowed" in text
    assert "_validation" not in text and "provider_warning" not in text.lower() and "zzz" not in text
    assert len(doc.tables) == 1 and doc.tables[0].rows[0].cells[1].text == "FY2023"


def test_advisory_warning_does_not_discard_sourced_data():
    data = {
        "provider_warning": "[Provider unavailable] RuntimeError: No configured provider returned a usable response",
        "sources": [{"url": "https://www.sec.gov/x", "title": "10-K", "period": "2025"}],
        "historical_financials": [{"fiscal_year": "FY2025", "revenue": 457_547_000, "source_url": "https://www.sec.gov/x"}],
    }
    brief = build_evidence_brief({"id": "d"}, [{"agent_type": "financial_analyst", "data": data, "success": True}])
    assert brief["successful_analysis_count"] == 1
    assert any(p["metric"] == "revenue" and p["source_id"] for p in brief["data_points"])
    failed = build_evidence_brief({"id": "d"}, [{"agent_type": "x", "data": {"error": "boom"}, "success": True}])
    assert failed["successful_analysis_count"] == 0


def test_unsourced_statement_figures_are_flagged_and_sourced_ones_are_not():
    doc = Document()
    render_payload(doc, {"annual_revenue": 457_547_000, "net_income": -12_000_000, "gross_margin": 0.7756},
                   verified={457_547_000.0})
    lines = {t.split(":")[0]: t for _, t in _doc_text(doc)}
    assert "unverified" not in lines["Annual Revenue"].lower()
    assert "unverified" in lines["Net Income"].lower()
    assert "unverified" not in lines["Gross Margin"].lower()  # derived ratio, not a statement line

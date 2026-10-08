from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest


def test_source_only_retrieval_confidence_is_not_misrepresented_as_calibrated():
    from app.core.reports.report_generator import _confidence_label

    assert _confidence_label({
        "confidence": 0.65,
        "data": {"confidence_basis": "not_calibrated_source_report"},
    }) == "Not calibrated"
    assert _confidence_label({
        "confidence": 0.65,
        "data": {"synthesis_status": "deterministic_source_report"},
    }) == "Not calibrated"
    assert _confidence_label({"confidence": 0.72, "data": {"confidence_basis": "model_assessed"}}) == "Not calibrated"
    assert _confidence_label({"confidence": 0.72, "data": {"confidence_calibrated": True}}) == "72%"
    assert _confidence_label({}) == "Not reported"


@pytest.mark.parametrize(
    ("fmt", "member"),
    [
        ("docx", "word/document.xml"),
        ("xlsx", "xl/workbook.xml"),
        ("pptx", "ppt/presentation.xml"),
    ],
)
def test_ooxml_artifact_validation_requires_expected_parts(fmt, member):
    from app.core.reports.report_guardrails import ReportGuardrails

    # An envelope with the right part names but no real content used to pass;
    # artifacts are now re-opened with their own library.
    out = BytesIO()
    with ZipFile(out, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr(member, "<root/>")

    assert ReportGuardrails.validate_artifact(fmt, out.getvalue())["valid"] is False
    assert ReportGuardrails.validate_artifact(fmt, b"not a document")["valid"] is False
    assert ReportGuardrails.validate_artifact(fmt, _real_artifact(fmt))["valid"] is True


def _real_artifact(fmt):
    if fmt == "docx":
        from docx import Document

        doc = Document()
        doc.add_paragraph("Real content")
        buf = BytesIO()
        doc.save(buf)
    elif fmt == "xlsx":
        from openpyxl import Workbook

        wb = Workbook()
        wb.active["A1"] = "Real content"
        buf = BytesIO()
        wb.save(buf)
    else:
        from pptx import Presentation

        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[6])
        buf = BytesIO()
        prs.save(buf)
    return buf.getvalue()


def test_pdf_validation_requires_header_and_eof_marker():
    from app.core.reports.report_guardrails import ReportGuardrails

    from reportlab.pdfgen import canvas

    buf = BytesIO()
    page = canvas.Canvas(buf)
    page.drawString(72, 720, "Real content")
    page.save()
    assert ReportGuardrails.validate_artifact("pdf", buf.getvalue())["valid"] is True
    assert ReportGuardrails.validate_artifact("pdf", b"%PDF-1.7\nbody\n%%EOF\n")["valid"] is False, "fake envelope"
    assert ReportGuardrails.validate_artifact("pdf", b"not a pdf")["valid"] is False


def test_sources_and_uses_refuses_to_invent_missing_company_inputs():
    from app.core.reports.report_guardrails import ReportGuardrails

    with pytest.raises(ValueError, match="explicit cash, EBITDA, and debt"):
        ReportGuardrails.generate_sources_and_uses({}, target_ev=100)


@pytest.mark.asyncio
async def test_cached_docx_download_returns_word_content_type(monkeypatch):
    from app import main
    from app.core.document_store import DocumentStore

    class FakeDocumentStore:
        async def get_document(self, deal_id, fmt):
            return b"docx-bytes" if fmt == "docx" else None

        async def get_document_meta(self, deal_id, fmt):
            return {
                "safe_filename": "Cedarline_Analytics",
                "release_status": "approved",
                "report_version": "v1",
                "analysis_fingerprint": "sha256:abc",
            }

    monkeypatch.setattr(DocumentStore, "get_instance", classmethod(lambda cls: FakeDocumentStore()))

    response = await main.download_deal_document("deal-123", "docx")

    assert response.status_code == 200
    assert response.media_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert response.headers["content-disposition"].endswith('DealForge_Cedarline_Analytics.docx"')
    assert response.body == b"docx-bytes"


@pytest.mark.asyncio
async def test_cached_document_download_is_blocked_until_approval(monkeypatch):
    from fastapi import HTTPException
    from app import main
    from app.core.document_store import DocumentStore

    class FakeDocumentStore:
        async def get_document(self, deal_id, fmt):
            return b"pending-review-report"

        async def get_document_meta(self, deal_id, fmt):
            return {
                "safe_filename": "Cedarline_Analytics",
                "release_status": "pending_review",
                "report_version": "v1",
                "analysis_fingerprint": "sha256:abc",
            }

    monkeypatch.setattr(DocumentStore, "get_instance", classmethod(lambda cls: FakeDocumentStore()))

    with pytest.raises(HTTPException) as exc:
        await main.download_deal_document("deal-123", "docx")
    assert exc.value.status_code == 409
    assert "approves" in exc.value.detail


@pytest.mark.asyncio
async def test_direct_legacy_report_generation_is_disabled():
    from fastapi import HTTPException
    from app.main import generate_deal_report

    with pytest.raises(HTTPException) as exc:
        await generate_deal_report("deal-123")
    assert exc.value.status_code == 409


@pytest.mark.asyncio
async def test_report_approval_records_reviewer_attestation(monkeypatch):
    from app import main
    from app.core.document_store import DocumentStore

    metadata = [
        {
            "format": fmt,
            "release_status": "pending_review",
            "report_version": "v1",
            "analysis_fingerprint": "sha256:abc",
        }
        for fmt in ("docx", "pdf", "pptx", "xlsx")
    ]

    class FakeDocumentStore:
        async def list_documents(self, deal_id):
            return [item.copy() for item in metadata]

        async def update_document_metadata(self, deal_id, fmt, updates):
            item = next(entry for entry in metadata if entry["format"] == fmt)
            item.update(updates)
            return item

    monkeypatch.setattr(DocumentStore, "get_instance", classmethod(lambda cls: FakeDocumentStore()))
    response = await main.approve_deal_documents(
        "deal-123", main.ReportApprovalRequest(reviewer="Reviewer A", attestation=True)
    )

    assert response["release_status"] == "approved"
    assert response["report_version"] == "v1"
    assert set(response["formats"]) == {"docx", "pdf", "pptx", "xlsx"}
    assert all(item["approved_by"] == "Reviewer A" for item in metadata)
    assert all(item["reviewer_attestation"] is True for item in metadata)


@pytest.mark.asyncio
async def test_report_approval_rejects_incomplete_bundle(monkeypatch):
    from fastapi import HTTPException
    from app import main
    from app.core.document_store import DocumentStore

    class FakeDocumentStore:
        async def list_documents(self, deal_id):
            return [{
                "format": "pdf",
                "report_version": "v1",
                "analysis_fingerprint": "sha256:abc",
            }]

    monkeypatch.setattr(DocumentStore, "get_instance", classmethod(lambda cls: FakeDocumentStore()))
    with pytest.raises(HTTPException) as exc:
        await main.approve_deal_documents(
            "deal-123", main.ReportApprovalRequest(reviewer="Reviewer A", attestation=True)
        )

    assert exc.value.status_code == 409
    assert "incomplete" in exc.value.detail


@pytest.mark.asyncio
async def test_report_generation_rejects_incomplete_persisted_task_lists(monkeypatch):
    from types import SimpleNamespace
    from fastapi import HTTPException
    from app import main

    class Store:
        async def get_deal(self, _deal_id):
            return {"id": "deal-1", "status": "completed"}

    class TaskManager:
        async def get_lists_for_deal(self, _deal_id):
            return [SimpleNamespace(items=[SimpleNamespace(status="pending", result=None)])]

    monkeypatch.setattr(main, "RedisStore", SimpleNamespace(get_instance=lambda: Store()))
    monkeypatch.setattr("app.core.tasks.task_manager.get_task_manager", lambda: TaskManager())

    with pytest.raises(HTTPException) as exc:
        await main.generate_deal_documents("deal-1")

    assert exc.value.status_code == 409
    assert "unfinished or unpersisted tasks" in exc.value.detail


def test_officecli_validation_runs_blocking_process_calls_off_event_loop(monkeypatch):
    import asyncio
    import subprocess
    import threading
    from app.core.reports.officecli_service import OfficeCLIService

    service = OfficeCLIService()
    monkeypatch.setattr(service, "is_available", lambda: True)
    caller_thread = threading.get_ident()
    worker_threads = []

    def fake_run(*args, **kwargs):
        worker_threads.append(threading.get_ident())
        stdout = "[]" if args[0] == "view" else "valid"
        return subprocess.CompletedProcess(args, 0, stdout=stdout, stderr="")

    monkeypatch.setattr(service, "_run_command", fake_run)
    result = asyncio.run(service.validate_document("test.docx"))

    assert result == {"success": True, "issues": []}
    assert len(worker_threads) == 2
    assert all(thread_id != caller_thread for thread_id in worker_threads)

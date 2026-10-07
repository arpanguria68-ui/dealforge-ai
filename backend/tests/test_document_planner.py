from app.core.reports.document_planner import (
    default_document_blueprint,
    latest_agent_outputs,
    validate_document_blueprint,
)


def test_latest_successful_agent_result_wins_over_old_and_new_failure():
    activities = [
        {"agent_type": "financial_analyst", "success": True, "data": {"value": 1}},
        {"agent_type": "financial_analyst", "success": False, "data": {"error": "retry failed"}},
        {"agent_type": "risk_assessor", "success": False, "data": {"error": "failed"}},
        {"agent_type": "financial_analyst", "success": True, "data": {"value": 2}},
    ]

    selected = latest_agent_outputs(activities)

    assert selected == [activities[-1], activities[2]]


def test_blueprint_cannot_add_sections_or_visuals_without_supporting_data():
    fallback = default_document_blueprint(
        {"target_company": "TestCo"},
        {"source_agents": [], "available_data": {"financial_metrics": False}, "unknowns": []},
    )
    candidate = {
        "sections": ["executive_summary", "investment_thesis", "valuation"],
        "visuals": ["ebitda_waterfall", "valuation_range"],
        "page_orientation": "portrait",
        "density": "compact",
    }

    result = validate_document_blueprint(candidate, fallback)

    assert result["sections"] == ["executive_summary"]
    assert result["visuals"] == []
    assert result["density"] == "compact"


def test_blueprint_layout_preferences_are_allowlisted():
    fallback = default_document_blueprint(
        {}, {"source_agents": [], "available_data": {}, "unknowns": []}
    )

    result = validate_document_blueprint(
        {"sections": ["financial_metrics"], "page_orientation": "diagonal", "density": "tiny"},
        fallback,
    )

    assert result["page_orientation"] == "portrait"
    assert result["density"] == "standard"


def test_docx_applies_validated_layout_and_includes_plan_and_qa_notes():
    from docx import Document
    from app.core.reports.report_generator import generate_docx
    from io import BytesIO

    content = generate_docx(
        {"target_company": "TestCo", "industry": "software"},
        {
            "_report_blueprint": {
                "audience": "Investment Committee",
                "sections": ["executive_summary", "diligence_gaps"],
                "planning_rationale": "No source-backed forecast was supplied.",
                "page_orientation": "landscape",
                "density": "compact",
                "omitted_sections": ["valuation"],
            },
            "_document_qa": {
                "status": "review_required",
                "warnings": ["Architect unavailable; deterministic defaults used."],
                "note": "Verify claims against cited records.",
            },
        },
        [],
    )
    document = Document(BytesIO(content))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert document.sections[0].page_width > document.sections[0].page_height
    assert "Evidence Coverage & Review Status" in text
    assert "Document Plan & Quality Notes" not in text
    assert "Omitted for insufficient recorded evidence" not in text
    assert "Architect unavailable" in text


def test_curated_payload_flags_cross_agent_metric_disagreements():
    from app.core.reports.document_planner import _metric_conflicts

    conflicts = _metric_conflicts([
        {"agent_type": "financial_analyst", "data": {"reported_revenue": 12}},
        {"agent_type": "market_researcher", "data": {"reported_revenue": 15}},
    ])

    assert len(conflicts) == 1
    assert {item["value"] for item in conflicts[0]["reports"]} == {"12", "15"}


def test_realistic_consulting_case_generates_valid_deliverables():
    import asyncio
    from io import BytesIO
    from docx import Document
    from app.agents.base import AgentOutput
    from app.core.reports.document_planner import prepare_document_payload
    from app.core.reports.report_generator import generate_docx, generate_pdf, generate_excel, generate_pptx
    from app.core.reports.report_guardrails import ReportGuardrails

    class FakeAgent:
        def __init__(self, data):
            self.data = data

        async def run(self, task, context=None):
            return AgentOutput(success=True, data=self.data, reasoning="fixture", confidence=0.8)

    class FakeRegistry:
        def get(self, name):
            if name == "business_analyst":
                return FakeAgent({"executive_summary": {
                    "situation": "Northstar Components supplies industrial customers.",
                    "complication": "Revenue differs across two recorded analyses.",
                    "question": "Can the reported operating profile be validated?",
                    "answer": "No recommendation recorded; reconcile reported revenue first.",
                }})
            if name == "report_architect":
                return FakeAgent({
                    "sections": ["executive_summary", "financial_metrics", "risk_assessment", "valuation"],
                    "visuals": ["valuation_range", "ebitda_waterfall"],
                    "omitted_sections": ["valuation"],
                    "density": "compact",
                    "page_orientation": "portrait",
                    "planning_rationale": "Focus the report on recorded evidence and open diligence.",
                })
            return None

    deal = {
        "id": "fictional-northstar-test", "name": "Northstar Components Diligence",
        "target_company": "Northstar Components", "industry": "industrial manufacturing",
        "status": "completed", "final_score": None,
    }
    activities = [
        {"agent_type": "financial_analyst", "success": True, "data": {"reported_revenue": 10.0}},
        {"agent_type": "financial_analyst", "success": True, "confidence": 0.82,
         "reasoning": "Management accounts report FY2025 revenue of $12.0m; audited statements not supplied.",
         "data": {"reported_revenue": 12.0, "ebitda_margin": "14%",
                  "key_findings": ["FY2025 revenue reported as $12.0m, unaudited"],
                  "data_gaps": ["Audited FY2025 statements not supplied"],
                  "citations": ["Fictional management accounts, p. 4"]}},
        {"agent_type": "market_researcher", "success": True,
         "reasoning": "Customer concentration could affect renewal durability.",
         "data": {"reported_revenue": 11.5,
                  "key_findings": ["Three customers represent a material share of sales"],
                  "citations": ["Fictional customer schedule, p. 2"]}},
        {"agent_type": "risk_assessor", "success": True,
         "reasoning": "Supplier concentration merits diligence.",
         "data": {"risks": [{"risk": "Single-source component exposure", "severity": "medium",
                              "mitigation": "Not provided", "evidence": "Supplier schedule"}],
                  "data_gaps": ["No supplier continuity plan supplied"]}},
    ]

    agent_results, analyst_data, evidence = asyncio.run(
        prepare_document_payload(FakeRegistry(), deal, activities)
    )
    assert len(agent_results) == 3
    assert agent_results[0]["data"]["reported_revenue"] == 12.0
    assert evidence["metric_conflicts"]
    assert "valuation" not in analyst_data["_report_blueprint"]["sections"]
    assert "valuation" in analyst_data["_report_blueprint"]["omitted_sections"]

    artifacts = {
        "docx": generate_docx(deal, analyst_data, agent_results),
        "pdf": generate_pdf(deal, analyst_data, agent_results),
        "xlsx": generate_excel(deal, analyst_data, agent_results),
        "pptx": generate_pptx(deal, analyst_data, agent_results),
    }
    for fmt, content in artifacts.items():
        assert ReportGuardrails.validate_artifact(fmt, content)["valid"]

    report = Document(BytesIO(artifacts["docx"]))
    doc_text = "\n".join(p.text for p in report.paragraphs)
    doc_text += "\n" + "\n".join(cell.text for table in report.tables for row in table.rows for cell in row.cells)
    assert "Metric reconciliation required" in doc_text
    assert "Audited FY2025 statements not supplied" in doc_text
    assert "Forecast status" not in doc_text

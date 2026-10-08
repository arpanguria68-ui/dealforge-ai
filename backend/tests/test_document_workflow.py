"""Adaptive document workflow: understand → plan → fill gaps → compose → render → review."""

import io
from types import SimpleNamespace

import pytest

from app.core.reports import document_workflow as dw


DEAL = {"id": "deal-doc", "name": "Project Acme", "target_company": "Acme SaaS", "industry": "software",
        "status": "completed", "current_stage": "deep_dive"}

FIN_RESULT = {
    "agent_type": "financial_analyst", "success": True,
    "data": {
        "summary": "Revenue grew steadily with improving margins.",
        "key_findings": ["Revenue grew 18% year over year", "Operating margin expanded"],
        "sources": [{"url": "https://sec.gov/acme-10k", "title": "Acme 10-K 2024"}],
        "historical_financials": [
            {"fiscal_year": "FY2023", "revenue": 100.0, "source_url": "https://sec.gov/acme-10k"},
            {"fiscal_year": "FY2024", "revenue": 118.0, "source_url": "https://sec.gov/acme-10k"},
        ],
        "valuation": {"enterprise_value": 950.0, "wacc": 0.095},
    },
}
RISK_RESULT = {
    "agent_type": "risk_assessor", "success": True,
    "data": {"risks": [
        {"risk": "Customer concentration", "severity": "high", "category": "Commercial", "mitigation": "Diversify top accounts"},
        {"risk": "Key person dependency", "severity": 3, "category": "Operational"},
    ]},
}
LEGAL_RESULT = {
    "agent_type": "legal_advisor", "success": True,
    "data": {"key_legal_risks": [{"risk": "Customer Concentration", "severity": "critical"}, "Pending IP litigation"]},
}


@pytest.fixture(autouse=True)
def _no_graph(monkeypatch):
    from app.core.knowledge_graph import service

    async def no_risks(deal_id, limit=10):
        return []

    monkeypatch.setattr(service, "risk_register", no_risks)


# ── Understand ──────────────────────────────────────────────────────────


@pytest.mark.parametrize("text,doc_type,formats,audience", [
    ("Prepare an IC memo for the board as PDF", "ic_memo", ["pdf"], "Board of Directors"),
    ("Need a one-pager teaser", "one_pager", ["docx", "pdf"], "Deal Team"),
    ("risk register in excel", "risk_report", ["xlsx"], "Risk Committee"),
    ("build the slides", "dd_report", ["pptx"], "Investment Committee"),
    ("", "dd_report", ["docx", "pdf", "pptx", "xlsx"], "Investment Committee"),
])
def test_request_interpretation(text, doc_type, formats, audience):
    intent = dw.interpret_request(dw.DocumentRequest(request=text))
    assert intent["doc_type"] == doc_type
    assert intent["formats"] == formats
    assert intent["audience"] == audience


def test_ambiguous_request_states_its_assumption_and_explicit_fields_win():
    intent = dw.interpret_request(dw.DocumentRequest(request="write something about this deal"))
    assert intent["doc_type"] == "ic_memo"
    assert any("assumed an Investment Committee memo" in a for a in intent["assumptions"])
    explicit = dw.interpret_request(dw.DocumentRequest(
        request="one pager", doc_type="risk_report", formats=["pdf", "bogus"], audience="Lenders"))
    assert (explicit["doc_type"], explicit["formats"], explicit["audience"]) == ("risk_report", ["pdf"], "Lenders")


# ── Inventory / risk register ───────────────────────────────────────────


def test_risk_register_merges_agents_and_graph_with_normalized_severity():
    graph = [{"name": "Pending IP litigation", "severity": 7, "category": "Legal", "description": "Patent suit"}]
    register = dw.collect_risk_register([RISK_RESULT, LEGAL_RESULT, {"agent_type": "x", "success": False,
                                         "data": {"risks": [{"risk": "ignored"}]}}], graph)
    names = [r["risk"] for r in register]
    assert names[0] == "Customer concentration"  # deduped, highest severity (critical=10) kept
    assert register[0]["severity"] == 10 and set(register[0]["sources"]) == {"risk_assessor", "legal_advisor"}
    ip = next(r for r in register if r["risk"] == "Pending IP litigation")
    assert set(ip["sources"]) == {"legal_advisor", "knowledge_graph"} and ip["severity"] == 7
    assert next(r for r in register if r["risk"] == "Key person dependency")["severity"] == 6  # 3/5 -> 6/10
    assert "ignored" not in names


# ── Plan ────────────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_plan_turns_missing_required_sections_into_gaps_with_owning_agents():
    plan, _, _ = await dw.build_plan(dw.DocumentRequest(request="IC memo"), DEAL, [FIN_RESULT])
    keys = plan.sections
    assert plan.doc_type == "ic_memo"
    assert "financial_metrics" in keys and "valuation" in keys
    assert plan.gaps == ["risk_assessment"]
    assert "risk_assessor" in plan.gap_agents
    assert "diligence_gaps" in keys
    assert any("fill_gaps=true" in q for q in plan.questions)
    assert "market_analysis" not in keys, "optional sections without evidence are omitted"


@pytest.mark.asyncio
async def test_architect_can_reorder_but_not_add_or_drop_required_sections():
    plan, _, _ = await dw.build_plan(dw.DocumentRequest(request="IC memo", use_architect=False),
                                     DEAL, [FIN_RESULT, RISK_RESULT])

    class Architect:
        async def run(self, task, context=None):
            assert context["allowed_sections"] == plan.sections
            return SimpleNamespace(success=True, data={"sections": [
                "risk_assessment", "hallucinated_section", "executive_summary", "valuation"]})

    registry = {"report_architect": Architect()}
    refined = await dw.refine_with_architect(plan, registry, DEAL)
    assert refined.sections[:3] == ["risk_assessment", "executive_summary", "valuation"]
    assert "hallucinated_section" not in refined.sections
    assert {"investment_thesis", "financial_metrics"} <= set(refined.sections)
    assert refined.planning_source == "report_architect_within_evidence_allowlist"


# ── Full run: gap filling, compose, render ──────────────────────────────


class _Agent:
    def __init__(self, data):
        self.data = data
        self.calls = 0

    async def run(self, task, context=None):
        self.calls += 1
        return SimpleNamespace(success=True, data=self.data)


@pytest.mark.asyncio
async def test_full_run_fills_gaps_and_renders_only_requested_formats():
    risk_agent = _Agent(RISK_RESULT["data"])
    registry = {"risk_assessor": risk_agent}
    out = await dw.run_document_workflow(
        dw.DocumentRequest(request="IC memo as word and pdf", fill_gaps=True, use_architect=False),
        DEAL, [FIN_RESULT], registry,
    )
    assert risk_agent.calls == 1
    assert out["gap_fill"][0]["agent"] == "risk_assessor"
    assert out["plan"]["gaps"] == [], "risk section covered after the gap agent ran"
    assert set(out["artifacts"]) == {"docx", "pdf"} and not out["errors"]

    from docx import Document

    text = "\n".join(p.text for p in Document(io.BytesIO(out["artifacts"]["docx"])).paragraphs)
    tables = Document(io.BytesIO(out["artifacts"]["docx"])).tables
    cells = {c.text for t in tables for row in t.rows for c in row.cells}
    assert "Investment Committee Memo" in text and "Risk Assessment" in text
    assert "Customer concentration" in cells and "Diversify top accounts" in cells
    assert out["artifacts"]["pdf"].startswith(b"%PDF")
    # Risk register also reaches legacy generators via analyst_data.
    assert out["analyst_data"]["risk_matrix"][0]["risk"] == "Customer concentration"


@pytest.mark.asyncio
async def test_risk_report_xlsx_has_register_sheet_and_review_banner():
    out = await dw.run_document_workflow(
        dw.DocumentRequest(request="risk report", formats=["xlsx"], use_architect=False),
        DEAL, [FIN_RESULT, RISK_RESULT, LEGAL_RESULT], None,
    )
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(out["artifacts"]["xlsx"]))
    assert "Risk Assessment" in wb.sheetnames
    rows = [[c.value for c in r] for r in wb["Risk Assessment"].iter_rows()]
    assert rows[0][:2] == ["Risk", "Severity"] and any(r[0] == "Pending IP litigation" for r in rows)
    assert out["model"]["review_status"] in {"review_required", "ready_for_review"}


def test_compose_never_emits_empty_sections_and_reports_them():
    plan = dw.DocumentPlan(doc_type="ic_memo", title="T", formats=["docx"], audience="IC", density="standard",
                           sections=["executive_summary", "market_analysis"], coverage={}, gaps=[], gap_agents=[])
    model = dw.compose_document_model(plan, DEAL, {"executive_summary": {"situation": "S"}},
                                      {"findings": [], "sources": []}, [], [])
    assert [s["key"] for s in model["sections"]] == ["executive_summary"]
    assert any("Market Analysis" in w for w in model["warnings"])


# ── API ─────────────────────────────────────────────────────────────────


class _FakeDocStore:
    def __init__(self):
        self.published = None

    async def replace_documents(self, deal_id, artifacts, metadata):
        self.published = (artifacts, metadata)
        return [{"format": f, **metadata} for f in artifacts]

    async def list_documents(self, deal_id):
        if not self.published:
            return []
        artifacts, metadata = self.published
        return [{"format": f, **metadata} for f in artifacts]


@pytest.fixture
def api(monkeypatch):
    from fastapi.testclient import TestClient

    import app.main as main
    from app.agents import base
    from app.core.document_store import DocumentStore

    async def load(deal_id):
        return DEAL, [], [], [FIN_RESULT, RISK_RESULT]

    store = _FakeDocStore()
    monkeypatch.setattr(main, "_load_report_inputs", load)
    monkeypatch.setattr(DocumentStore, "get_instance", classmethod(lambda cls: store))
    monkeypatch.setattr(base, "get_agent_registry", lambda: None)
    return TestClient(main.app), store


def test_plan_endpoint_is_a_dry_run(api):
    client, store = api
    res = client.post("/api/v1/deals/deal-doc/documents/plan", json={"request": "one pager for the board"})
    assert res.status_code == 200
    plan = res.json()["plan"]
    assert plan["doc_type"] == "one_pager" and plan["audience"] == "Board of Directors"
    assert [s["key"] for s in plan["sections"]][0] == "snapshot"
    assert store.published is None


def test_generate_adaptive_publishes_planned_formats_for_review(api, monkeypatch):
    client, store = api
    res = client.post("/api/v1/deals/deal-doc/documents/generate",
                      json={"request": "risk report as pdf", "use_architect": False})
    body = res.json()
    assert res.status_code == 200 and body["status"] == "complete", body
    artifacts, metadata = store.published
    assert set(artifacts) == {"pdf"}
    assert metadata["doc_type"] == "risk_report" and metadata["expected_formats"] == ["pdf"]
    assert metadata["release_status"] == "pending_review"

    import app.main as main

    assert main._expected_report_formats(store_docs := [{"expected_formats": ["pdf"]}]) == {"pdf"}
    assert main._expected_report_formats([{}]) == main.REQUIRED_REPORT_FORMATS


def test_legacy_pack_honours_requested_formats(api, monkeypatch):
    client, store = api
    from app.core import provenance

    class _Prov:
        async def get_records(self, deal_id):
            return []

    monkeypatch.setattr(provenance, "get_provenance_collector", lambda: _Prov())
    res = client.post("/api/v1/deals/deal-doc/documents/generate", json={"request": "build the slides"})
    body = res.json()
    assert res.status_code == 200 and body["status"] == "complete", body
    artifacts, metadata = store.published
    assert set(artifacts) == {"pptx"} and metadata["expected_formats"] == ["pptx"]
    assert metadata["doc_type"] == "dd_report"


def test_officecli_clean_report_is_not_treated_as_issues():
    from app.core.reports.officecli_service import OfficeCLIService

    norm = OfficeCLIService._normalize_issues
    assert norm({"success": True, "data": {"count": 0, "issues": []}}) == []
    assert norm({"data": {"issues": [{"severity": "warning", "msg": "style"}]}}) == []
    assert len(norm({"data": {"issues": [{"severity": "error", "msg": "broken xml"}]}})) == 1
    assert len(norm([{"msg": "unknown level counts as error"}])) == 1
    assert norm("garbage") == []


def test_numeric_risk_severity_renders_in_legacy_generators():
    from app.core.reports.report_generator import _severity_label

    assert [_severity_label(v) for v in (3, 5, 8, 10, "High", None)] == ["Medium", "Critical", "High", "Critical", "High", "Not rated"]

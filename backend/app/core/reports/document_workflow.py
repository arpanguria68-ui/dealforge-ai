"""Adaptive, evidence-bounded document workflow (understand → plan → fill → compose → render → review).

The Reports Hub used to produce one fixed deliverable: a "Due Diligence
Report" in all four formats, with a blueprint the generators mostly ignored,
no use of the ReportArchitect agent, and no connection to the cross-agent
risk register. This module adapts the deliverable to the request, the way a
human analyst (or Cowork) would:

1. Understand: infer document type, formats and audience from the request
   text. Assumptions are stated, not silent.
2. Inventory: map every candidate section to the agents and evidence that
   feed it (covered / partial / missing).
3. Plan: keep the sections the evidence supports, turn missing required
   sections into explicit gaps, and let ReportArchitect reorder or trim
   within that allow-list.
4. Fill gaps (opt-in, bounded): dispatch the agents that own missing
   sections and fold their outputs into the evidence.
5. Compose: build a document model only from curated, cited data. No new
   model prose, so the existing hallucination guarantees hold.
6. Render: only the requested formats, concurrently.
7. Review: QA status plus concrete next actions.

The legacy full "dd_report" path is unchanged and still renders through
report_generator.
"""

from __future__ import annotations

import asyncio
import io
import re
import uuid
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional, Tuple

import structlog

logger = structlog.get_logger(__name__)

SUPPORTED_FORMATS = ("docx", "pdf", "xlsx", "pptx")
ADAPTIVE_RENDERERS = ("docx", "pdf", "xlsx", "pptx")

SECTION_TITLES = {
    "snapshot": "Deal Snapshot",
    "executive_summary": "Executive Summary",
    "investment_thesis": "Investment Thesis",
    "financial_metrics": "Financial Metrics",
    "valuation": "Valuation",
    "market_analysis": "Market Analysis",
    "risk_assessment": "Risk Assessment",
    "diligence_gaps": "Diligence Gaps",
    "next_steps": "Next Steps",
    "sources_methodology": "Sources and Methodology",
}

# Which agents own each section, so a gap can be filled by the right one.
SECTION_AGENTS = {
    "financial_metrics": ["financial_analyst"],
    "valuation": ["dcf_lbo_architect", "financial_analyst"],
    "market_analysis": ["market_researcher"],
    "risk_assessment": ["risk_assessor", "legal_advisor"],
    "investment_thesis": ["market_researcher"],
}

DOC_TYPES: Dict[str, Dict[str, Any]] = {
    "dd_report": {
        "title": "Due Diligence Report", "formats": ["docx", "pdf", "pptx", "xlsx"],
        "audience": "Investment Committee", "density": "standard", "legacy": True,
        "required": ["executive_summary"],
        "optional": ["financial_metrics", "market_analysis", "risk_assessment", "valuation",
                     "diligence_gaps", "sources_methodology", "next_steps"],
    },
    "ic_memo": {
        "title": "Investment Committee Memo", "formats": ["docx", "pdf"],
        "audience": "Investment Committee", "density": "standard",
        "required": ["executive_summary", "investment_thesis", "financial_metrics", "risk_assessment"],
        "optional": ["valuation", "market_analysis", "diligence_gaps", "next_steps", "sources_methodology"],
    },
    "one_pager": {
        "title": "Deal One-Pager", "formats": ["docx", "pdf"],
        "audience": "Deal Team", "density": "compact", "item_limit": 5,
        "required": ["snapshot"],
        "optional": ["financial_metrics", "risk_assessment", "diligence_gaps"],
    },
    "risk_report": {
        "title": "Risk Assessment Report", "formats": ["docx", "xlsx"],
        "audience": "Risk Committee", "density": "detailed",
        "required": ["risk_assessment"],
        "optional": ["executive_summary", "diligence_gaps", "next_steps", "sources_methodology"],
    },
    "board_deck": {
        "title": "Board Deck", "formats": ["pptx", "pdf"],
        "audience": "Board of Directors", "density": "compact", "item_limit": 6,
        "required": ["executive_summary", "risk_assessment"],
        "optional": ["snapshot", "investment_thesis", "financial_metrics", "valuation", "diligence_gaps", "next_steps"],
    },
    "financial_summary": {
        "title": "Financial Summary", "formats": ["docx", "xlsx"],
        "audience": "Deal Team", "density": "standard",
        "required": ["financial_metrics"],
        "optional": ["executive_summary", "valuation", "diligence_gaps", "sources_methodology"],
    },
}

_TYPE_KEYWORDS = [
    ("one_pager", ("one pager", "one-pager", "onepager", "teaser", "snapshot", "single page", "1-pager", "tear sheet")),
    ("risk_report", ("risk report", "risk register", "risk assessment", "red flag", "risk memo")),
    ("financial_summary", ("financial summary", "financials pack", "financial pack", "financial overview", "model summary")),
    ("dd_report", ("due diligence report", "dd report", "full report", "complete report", "full pack", "all formats")),
    ("board_deck", ("board deck", "deck", "slides", "presentation", "powerpoint", "pptx")),
    ("ic_memo", ("ic memo", "investment memo", "investment committee", "investment memorandum", "memo")),
]
_FORMAT_KEYWORDS = [
    ("pptx", ("deck", "slides", "powerpoint", "pptx", "presentation")),
    ("xlsx", ("excel", "xlsx", "spreadsheet", "workbook")),
    ("pdf", ("pdf",)),
    ("docx", ("word", "docx", ".doc ", " doc ")),
]
_AUDIENCE_KEYWORDS = [
    ("Board of Directors", ("board",)),
    ("Limited Partners", (" lp", "lps", "limited partner", "investors")),
    ("Investment Committee", (" ic ", "ic memo", "investment committee")),
    ("Management Team", ("management", "leadership team")),
    ("Lenders", ("lender", "bank", "credit committee")),
]


@dataclass
class DocumentRequest:
    request: str = ""
    doc_type: Optional[str] = None
    formats: Optional[List[str]] = None
    audience: Optional[str] = None
    fill_gaps: bool = False
    max_gap_agents: int = 2
    use_architect: bool = True


@dataclass
class DocumentPlan:
    doc_type: str
    title: str
    formats: List[str]
    audience: str
    density: str
    sections: List[str]
    coverage: Dict[str, Dict[str, Any]]
    gaps: List[str]
    gap_agents: List[str]
    assumptions: List[str] = field(default_factory=list)
    questions: List[str] = field(default_factory=list)
    planning_source: str = "evidence_coverage"
    legacy: bool = False

    def as_dict(self) -> Dict[str, Any]:
        return {
            "doc_type": self.doc_type, "title": self.title, "formats": self.formats,
            "audience": self.audience, "density": self.density,
            "sections": [{"key": s, "title": SECTION_TITLES.get(s, s)} for s in self.sections],
            "coverage": self.coverage, "gaps": self.gaps, "gap_agents": self.gap_agents,
            "assumptions": self.assumptions, "questions": self.questions,
            "planning_source": self.planning_source, "legacy": self.legacy,
        }


# ── 1. Understand ──────────────────────────────────────────────────────

def interpret_request(req: DocumentRequest) -> Dict[str, Any]:
    """Infer type/formats/audience from text; explicit fields always win."""
    text = f" {(req.request or '').lower()} "
    assumptions: List[str] = []
    questions: List[str] = []

    doc_type = req.doc_type if req.doc_type in DOC_TYPES else None
    if req.doc_type and not doc_type:
        assumptions.append(f"Unknown document type '{req.doc_type}'; inferred from the request instead.")
    if not doc_type:
        doc_type = next((t for t, words in _TYPE_KEYWORDS if any(w in text for w in words)), None)
    requested_formats = [f for f in (req.formats or []) if f in SUPPORTED_FORMATS]
    text_formats = [f for f, words in _FORMAT_KEYWORDS if any(w in text for w in words)]
    if not doc_type:
        if text.strip() == "":
            doc_type = "dd_report"  # no request: legacy full pack (backwards compatible)
        elif "pptx" in (requested_formats or text_formats):
            doc_type = "board_deck"
            assumptions.append("A deck was requested; using the board deck layout.")
        else:
            doc_type = "ic_memo"
            assumptions.append(
                "Document type not specified; assumed an Investment Committee memo. Ask for a "
                "'one-pager', 'risk report', 'financial summary' or 'full report' to change it."
            )

    spec = DOC_TYPES[doc_type]
    formats = requested_formats or text_formats or list(spec["formats"])
    formats = list(dict.fromkeys(formats))
    if not (requested_formats or text_formats):
        assumptions.append(f"Formats not specified; using {', '.join(formats)} for a {spec['title'].lower()}.")

    audience = req.audience or next(
        (label for label, words in _AUDIENCE_KEYWORDS if any(w in text for w in words)), None
    )
    if not audience:
        audience = spec["audience"]
    return {"doc_type": doc_type, "formats": formats, "audience": audience,
            "assumptions": assumptions, "questions": questions}


_DOC_VERBS = ("create", "generate", "prepare", "draft", "write", "make", "build", "produce",
              "export", "give me", "send me", "put together", "compile", "i need", "need a")
_DOC_NOUNS = ("memo", "report", "one-pager", "one pager", "onepager", "teaser", "deck", "slides",
              "presentation", "tear sheet", "pdf", "word doc", "docx", "spreadsheet", "excel", "deliverable")


def detect_document_request(text: str) -> bool:
    """True when a chat message asks for a deliverable (verb + document noun)."""
    lower = f" {(text or '').lower()} "
    return any(v in lower for v in _DOC_VERBS) and any(n in lower for n in _DOC_NOUNS)


# ── 2. Inventory ───────────────────────────────────────────────────────

_RISK_FIELDS = ("risks", "top_risks", "key_risks", "key_legal_risks", "financial_risks", "flags", "red_flags")
_SEVERITY_WORDS = {"low": 3, "medium": 5, "moderate": 5, "high": 8, "critical": 10, "severe": 9}


def _severity_10(value: Any) -> Optional[float]:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value) * 2 if value <= 5 else float(min(value, 10))
    if isinstance(value, str):
        word = value.strip().lower()
        if word in _SEVERITY_WORDS:
            return float(_SEVERITY_WORDS[word])
        try:
            return _severity_10(float(word))
        except ValueError:
            return None
    return None


def collect_risk_register(agent_results: List[Dict[str, Any]], graph_risks: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """One deduplicated risk register from every agent plus the knowledge graph."""
    register: Dict[str, Dict[str, Any]] = {}

    def add(name: Any, severity: Any, category: Any, source: str, mitigation: Any = None, evidence: Any = None):
        title = str(name or "").strip()
        if not title:
            return
        key = re.sub(r"[^a-z0-9]+", " ", title.lower()).strip()[:120]
        sev = _severity_10(severity)
        current = register.get(key)
        if current is None:
            register[key] = {
                "risk": title[:240], "severity": sev, "category": str(category or "General")[:80],
                "sources": [source], "mitigation": (str(mitigation)[:300] if mitigation else None),
                "evidence": (str(evidence)[:300] if evidence else None),
            }
        else:
            if source not in current["sources"]:
                current["sources"].append(source)
            if sev is not None and (current["severity"] is None or sev > current["severity"]):
                current["severity"] = sev
            current["mitigation"] = current["mitigation"] or (str(mitigation)[:300] if mitigation else None)

    for result in agent_results:
        data = result.get("data") if isinstance(result.get("data"), dict) else {}
        agent = str(result.get("agent_type") or "unknown")
        if result.get("success") is False:
            continue
        for fieldname in _RISK_FIELDS:
            items = data.get(fieldname)
            for item in items if isinstance(items, list) else []:
                if isinstance(item, dict):
                    add(item.get("risk") or item.get("name") or item.get("title") or item.get("description"),
                        item.get("severity") or item.get("level") or item.get("rating"),
                        item.get("category") or item.get("type"), agent,
                        item.get("mitigation"), item.get("evidence") or item.get("rationale"))
                elif isinstance(item, str):
                    add(item, None, None, agent)
    for risk in graph_risks:
        add(risk.get("name"), risk.get("severity"), risk.get("category"), "knowledge_graph",
            evidence=risk.get("description"))
    ordered = sorted(register.values(), key=lambda r: (r["severity"] is None, -(r["severity"] or 0)))
    return ordered[:40]


def _valuation_points(agent_results: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    terms = ("valuation", "enterprise_value", "equity_value", "dcf", "irr", "moic", "wacc", "terminal", "multiple", "price_target")
    points = []
    for result in agent_results:
        if result.get("success") is False:
            continue
        data = result.get("data") if isinstance(result.get("data"), dict) else {}
        agent = str(result.get("agent_type") or "unknown")
        stack: List[Tuple[Dict[str, Any], str]] = [(data, "")]
        while stack and len(points) < 30:
            current, prefix = stack.pop()
            for name, value in current.items():
                path = f"{prefix}.{name}" if prefix else str(name)
                if isinstance(value, dict) and len(path) < 80:
                    stack.append((value, path))
                elif (isinstance(value, (int, float)) and not isinstance(value, bool)
                      and any(t in path.lower() for t in terms)):
                    points.append({"metric": " › ".join(part.replace("_", " ").title() for part in path.split(".")),
                                   "value": value, "agent": agent})
    return points


def _agent_succeeded(agent_results: List[Dict[str, Any]], names: Tuple[str, ...]) -> bool:
    return any(
        str(r.get("agent_type", "")) in names and r.get("success") is not False and isinstance(r.get("data"), dict)
        for r in agent_results
    )


def build_inventory(
    deal: Dict[str, Any], agent_results: List[Dict[str, Any]], evidence: Dict[str, Any],
    risk_register: List[Dict[str, Any]],
) -> Dict[str, Dict[str, Any]]:
    """Section → coverage status and the agents/evidence feeding it."""
    findings = evidence.get("findings", [])
    data_points = evidence.get("data_points", [])
    analyses = evidence.get("successful_analysis_count", 0)
    valuation = _valuation_points(agent_results)

    def entry(status: str, feeders: List[str], detail: str) -> Dict[str, Any]:
        return {"status": status, "feeding_agents": sorted(set(feeders)), "detail": detail}

    market_agents = [str(r.get("agent_type")) for r in agent_results if "market" in str(r.get("agent_type", ""))]
    return {
        "snapshot": entry("covered" if deal.get("target_company") or deal.get("name") else "partial", [],
                          "Deal identity and recorded status."),
        "executive_summary": entry("covered" if analyses else "missing", [f.get("agent", "") for f in findings],
                                   f"{analyses} successful analyses."),
        "investment_thesis": entry("covered" if len(findings) >= 2 else ("partial" if findings else "missing"),
                                   [f.get("agent", "") for f in findings], f"{len(findings)} recorded findings."),
        "financial_metrics": entry("covered" if data_points else "missing", [p.get("agent", "") for p in data_points],
                                   f"{len(data_points)} financial data points."),
        "valuation": entry("covered" if valuation else "missing", [p["agent"] for p in valuation],
                           f"{len(valuation)} valuation figures."),
        "market_analysis": entry(
            "covered" if _agent_succeeded(agent_results, tuple(market_agents) or ("market_researcher",)) else "missing",
            market_agents, "Market research output."),
        "risk_assessment": entry("covered" if risk_register else "missing",
                                 [s for r in risk_register for s in r["sources"]],
                                 f"{len(risk_register)} risks across agents and the knowledge graph."),
        "diligence_gaps": entry("covered" if evidence.get("unknowns") else "missing", [],
                                f"{len(evidence.get('unknowns', []))} recorded unknowns."),
        "next_steps": entry("covered" if evidence.get("unknowns") else "missing", [], "Derived from open diligence items."),
        "sources_methodology": entry("covered" if evidence.get("sources") else "partial", [],
                                     f"{len(evidence.get('sources', []))} cited sources."),
    }


# ── 3. Plan ────────────────────────────────────────────────────────────

def plan_document(
    req: DocumentRequest, deal: Dict[str, Any], agent_results: List[Dict[str, Any]],
    evidence: Dict[str, Any], risk_register: List[Dict[str, Any]],
) -> DocumentPlan:
    intent = interpret_request(req)
    spec = DOC_TYPES[intent["doc_type"]]
    coverage = build_inventory(deal, agent_results, evidence, risk_register)
    sections, gaps = [], []
    for key in spec["required"]:
        if coverage[key]["status"] == "missing":
            gaps.append(key)
        else:
            sections.append(key)
    for key in spec["optional"]:
        if coverage[key]["status"] != "missing":
            sections.append(key)
    if gaps and "diligence_gaps" not in sections:
        sections.append("diligence_gaps")
    # Narrative order (snapshot → summary → thesis → numbers → risks → gaps →
    # next steps → sources); ReportArchitect may still reorder within it.
    canonical = list(SECTION_TITLES)
    sections.sort(key=canonical.index)
    gap_agents = list(dict.fromkeys(a for g in gaps for a in SECTION_AGENTS.get(g, [])))
    questions = list(intent["questions"])
    if gaps:
        titles = ", ".join(SECTION_TITLES[g] for g in gaps)
        questions.append(
            f"Required sections without evidence: {titles}. Generate with fill_gaps=true to run "
            f"{', '.join(gap_agents) or 'the owning agents'}, or accept them as listed diligence gaps."
        )
    if not evidence.get("successful_analysis_count", 0):
        questions.append("No successful analyses are recorded for this deal; run the analysis first.")
    return DocumentPlan(
        doc_type=intent["doc_type"], title=spec["title"], formats=intent["formats"],
        audience=intent["audience"], density=spec["density"], sections=sections, coverage=coverage,
        gaps=gaps, gap_agents=gap_agents, assumptions=intent["assumptions"], questions=questions,
        legacy=bool(spec.get("legacy")),
    )


async def refine_with_architect(plan: DocumentPlan, registry: Any, deal: Dict[str, Any], timeout: float = 30.0) -> DocumentPlan:
    """Let ReportArchitect reorder/trim optional sections; never add or drop required ones."""
    agent = registry.get("report_architect") if registry is not None else None
    if agent is None or plan.legacy:
        return plan
    required = set(DOC_TYPES[plan.doc_type]["required"]) - set(plan.gaps)
    try:
        result = await asyncio.wait_for(agent.run(
            f"Plan the section order for a {plan.title} for {plan.audience}.",
            context={
                "deal_id": deal.get("id"), "company_name": deal.get("target_company") or deal.get("name"),
                "industry": deal.get("industry"), "audience": plan.audience,
                "allowed_sections": plan.sections, "available_visuals": [],
                "available_data": {k: v["status"] for k, v in plan.coverage.items() if k in plan.sections},
            },
        ), timeout=timeout)
    except Exception as exc:
        logger.info("report_architect_unavailable", error_type=type(exc).__name__)
        return plan
    proposed = getattr(result, "data", {}).get("sections") if getattr(result, "success", False) else None
    if not isinstance(proposed, list):
        return plan
    ordered = [s for s in dict.fromkeys(str(p) for p in proposed) if s in plan.sections]
    ordered += [s for s in plan.sections if s in required and s not in ordered]
    if "diligence_gaps" in plan.sections and plan.gaps and "diligence_gaps" not in ordered:
        ordered.append("diligence_gaps")
    if ordered:
        plan.sections = ordered
        plan.planning_source = "report_architect_within_evidence_allowlist"
    return plan


# ── 4. Fill gaps ───────────────────────────────────────────────────────

async def fill_gaps(
    plan: DocumentPlan, deal: Dict[str, Any], registry: Any, *, max_agents: int = 2, timeout: float = 120.0,
) -> List[Dict[str, Any]]:
    """Run the agents that own missing required sections (bounded, concurrent)."""
    if registry is None or not plan.gap_agents:
        return []
    names = [n for n in plan.gap_agents if registry.get(n) is not None][: max(0, max_agents)]
    target = deal.get("target_company") or deal.get("name") or "the target"

    async def run_one(name: str) -> Optional[Dict[str, Any]]:
        sections = [SECTION_TITLES[g] for g in plan.gaps if name in SECTION_AGENTS.get(g, [])]
        context = {"deal_id": deal.get("id"), "target_company": target, "industry": deal.get("industry"),
                   "deal_name": deal.get("name")}
        agent = registry.get(name)
        try:
            agent._current_context = context
        except Exception:
            pass
        try:
            out = await asyncio.wait_for(agent.run(
                f"Provide the {', '.join(sections)} needed for a {plan.title} on {target}. "
                "Use only cited evidence and list material unknowns explicitly.",
                context=context,
            ), timeout=timeout)
        except Exception as exc:
            logger.warning("document_gap_agent_failed", agent=name, error_type=type(exc).__name__)
            return None
        if not getattr(out, "success", False) or not isinstance(getattr(out, "data", None), dict):
            return None
        return {"agent_type": name, "task_id": f"docgap-{uuid.uuid4().hex[:8]}", "success": True,
                "data": out.data, "source": "document_gap_fill"}

    results = await asyncio.gather(*(run_one(n) for n in names))
    return [r for r in results if r]


# ── 5. Compose (no new model prose) ────────────────────────────────────

def _fmt_value(value: Any) -> str:
    if isinstance(value, bool) or value is None:
        return str(value)
    if isinstance(value, (int, float)):
        if abs(value) >= 1_000_000:
            return f"{value:,.0f}"
        if abs(value) >= 1 or value == 0:
            return f"{value:,.2f}".rstrip("0").rstrip(".")
        # Rates and ratios (WACC 0.095) must not be rounded to 0.1.
        return f"{value:.4g}"
    return str(value)


def compose_document_model(
    plan: DocumentPlan, deal: Dict[str, Any], analyst_data: Dict[str, Any], evidence: Dict[str, Any],
    agent_results: List[Dict[str, Any]], risk_register: List[Dict[str, Any]],
) -> Dict[str, Any]:
    limit = DOC_TYPES[plan.doc_type].get("item_limit") or (8 if plan.density == "compact" else 25)
    summary = analyst_data.get("executive_summary") or {}
    qa = analyst_data.get("_document_qa") or {}
    sections: List[Dict[str, Any]] = []

    def section(key: str, blocks: List[Dict[str, Any]]):
        blocks = [b for b in blocks if b.get("text") or b.get("items") or b.get("rows")]
        if blocks:
            sections.append({"key": key, "title": SECTION_TITLES[key], "blocks": blocks})

    for key in plan.sections:
        if key == "snapshot":
            section(key, [{"type": "table", "columns": ["Item", "Value"], "rows": [
                ["Target", deal.get("target_company") or deal.get("name") or "Not recorded"],
                ["Industry", deal.get("industry") or "Not recorded"],
                ["Stage", deal.get("current_stage") or deal.get("status") or "Not recorded"],
                ["Recorded recommendation", evidence.get("recommendation") or "None recorded; human review required"],
                ["Analyses completed", str(evidence.get("successful_analysis_count", 0))],
                ["Cited sources", str(len(evidence.get("sources", [])))],
            ]}])
        elif key == "executive_summary":
            section(key, [
                {"type": "paragraph", "text": summary.get("situation", "")},
                {"type": "paragraph", "text": summary.get("complication", "")},
                {"type": "paragraph", "label": "Decision question", "text": summary.get("question", "")},
                {"type": "paragraph", "label": "Recorded answer", "text": summary.get("answer", "")},
            ])
        elif key == "investment_thesis":
            takeaways = analyst_data.get("key_takeaways") or []
            items = [f"{t.get('title')}: {t.get('description')}" for t in takeaways if isinstance(t, dict)]
            if not items:
                items = [f"{f.get('agent', '').replace('_', ' ').title()}: {f.get('text')}"
                         for f in evidence.get("findings", [])[:limit]]
            section(key, [{"type": "bullets", "items": items[:limit]}])
        elif key == "financial_metrics":
            rows = [[str(p.get("metric", "")).replace("_", " ").title(), str(p.get("period") or "n/a"),
                     _fmt_value(p.get("value")), p.get("source_id") or "not cited", str(p.get("basis") or "")]
                    for p in evidence.get("data_points", [])[: limit * 2]]
            section(key, [{"type": "table", "columns": ["Metric", "Period", "Value", "Source", "Basis"], "rows": rows}])
        elif key == "valuation":
            rows = [[p["metric"], _fmt_value(p["value"]), p["agent"]] for p in _valuation_points(agent_results)[:limit]]
            section(key, [
                {"type": "table", "columns": ["Measure", "Value", "Agent"], "rows": rows},
                {"type": "paragraph", "text": "Valuation figures are reproduced from agent outputs and are not independently verified."},
            ])
        elif key == "market_analysis":
            items = [f.get("text") for f in evidence.get("findings", []) if "market" in str(f.get("agent", ""))]
            section(key, [{"type": "bullets", "items": [i for i in items if i][:limit]}])
        elif key == "risk_assessment":
            rows = [[r["risk"], "n/a" if r["severity"] is None else f"{r['severity']:.0f}/10", r["category"],
                     ", ".join(r["sources"]), r.get("mitigation") or "Not recorded"] for r in risk_register[:limit * 2]]
            section(key, [
                {"type": "table", "columns": ["Risk", "Severity", "Category", "Raised by", "Mitigation"], "rows": rows},
                {"type": "paragraph", "text": "Combined register from agent outputs, red-team flags and the deal knowledge graph; severities are agent-assessed."},
            ])
        elif key == "diligence_gaps":
            items = [u.get("text") if isinstance(u, dict) else str(u) for u in evidence.get("unknowns", [])]
            items += [f"No evidence recorded for {SECTION_TITLES[g]}." for g in plan.gaps]
            section(key, [{"type": "bullets", "items": [i for i in items if i][: limit * 2]}])
        elif key == "next_steps":
            section(key, [{"type": "bullets", "items": [a for a in analyst_data.get("action_items", []) if a][:limit]}])
        elif key == "sources_methodology":
            items = [f"[{s.get('id')}] {s.get('title') or s.get('url')}" + (f" ({s.get('url')})" if s.get("title") else "")
                     for s in evidence.get("sources", [])[:40]]
            section(key, [
                {"type": "bullets", "items": items},
                {"type": "paragraph", "text": qa.get("note") or "AI-generated synthesis; verify material statements against cited sources."},
            ])

    rendered = {s["key"] for s in sections}
    empty = [k for k in plan.sections if k not in rendered]
    return {
        "title": plan.title,
        "subtitle": deal.get("target_company") or deal.get("name") or "",
        "audience": plan.audience,
        "density": plan.density,
        "review_status": qa.get("status", "review_required"),
        "warnings": list(qa.get("warnings", [])) + [f"Planned section had no content: {SECTION_TITLES[k]}" for k in empty],
        "sections": sections,
    }


# ── 6. Render ──────────────────────────────────────────────────────────

def render_docx(model: Dict[str, Any]) -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10 if model["density"] == "compact" else 11)
    doc.add_heading(model["title"], level=0)
    doc.add_paragraph(f"{model['subtitle']} | Prepared for: {model['audience']}")
    if model["review_status"] != "ready_for_review" or model["warnings"]:
        p = doc.add_paragraph()
        p.add_run("Review required before distribution. ").bold = True
        p.add_run("; ".join(model["warnings"][:4]) or "AI-generated content; verify against sources.")
    for sec in model["sections"]:
        doc.add_heading(sec["title"], level=1)
        for block in sec["blocks"]:
            if block["type"] == "paragraph":
                para = doc.add_paragraph()
                if block.get("label"):
                    para.add_run(f"{block['label']}: ").bold = True
                para.add_run(block["text"])
            elif block["type"] == "bullets":
                for item in block["items"]:
                    doc.add_paragraph(str(item), style="List Bullet")
            elif block["type"] == "table":
                table = doc.add_table(rows=1, cols=len(block["columns"]))
                table.style = "Table Grid"
                for cell, name in zip(table.rows[0].cells, block["columns"]):
                    cell.text = name
                    for run in cell.paragraphs[0].runs:
                        run.bold = True
                for row in block["rows"]:
                    cells = table.add_row().cells
                    for cell, value in zip(cells, row):
                        cell.text = str(value)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render_pdf(model: Dict[str, Any]) -> bytes:
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    body = styles["BodyText"]
    story = [Paragraph(escape(model["title"]), styles["Title"]),
             Paragraph(escape(f"{model['subtitle']} | Prepared for: {model['audience']}"), body), Spacer(1, 8)]
    if model["review_status"] != "ready_for_review" or model["warnings"]:
        story.append(Paragraph("<b>Review required before distribution.</b> "
                               + escape("; ".join(model["warnings"][:4]) or "Verify against sources."), body))
    for sec in model["sections"]:
        story.append(Paragraph(escape(sec["title"]), styles["Heading2"]))
        for block in sec["blocks"]:
            if block["type"] == "paragraph":
                label = f"<b>{escape(block['label'])}:</b> " if block.get("label") else ""
                story.append(Paragraph(label + escape(block["text"]), body))
            elif block["type"] == "bullets":
                for item in block["items"]:
                    story.append(Paragraph("• " + escape(str(item)), body))
            elif block["type"] == "table":
                data = [[Paragraph(f"<b>{escape(c)}</b>", body) for c in block["columns"]]]
                data += [[Paragraph(escape(str(v)), body) for v in row] for row in block["rows"]]
                table = Table(data, repeatRows=1)
                table.setStyle(TableStyle([
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF5")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ]))
                story.append(table)
            story.append(Spacer(1, 4))
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=letter, title=model["title"]).build(story)
    return buf.getvalue()


def render_xlsx(model: Dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.append([model["title"]])
    summary["A1"].font = Font(bold=True, size=14)
    summary.append([model["subtitle"], f"Prepared for: {model['audience']}", f"Review: {model['review_status']}"])
    for warning in model["warnings"][:10]:
        summary.append(["Warning", warning])
    used = {"Summary"}
    for sec in model["sections"]:
        for block in sec["blocks"]:
            if block["type"] == "table":
                name = re.sub(r"[\[\]:*?/\\]", "", sec["title"])[:31] or "Sheet"
                base, n = name, 2
                while name in used:
                    name = f"{base[:28]} {n}"
                    n += 1
                used.add(name)
                ws = wb.create_sheet(name)
                ws.append(block["columns"])
                for cell in ws[1]:
                    cell.font = Font(bold=True)
                for row in block["rows"]:
                    ws.append([str(v) for v in row])
            elif block["type"] in ("paragraph", "bullets"):
                summary.append([])
                summary.append([sec["title"]])
                summary.cell(row=summary.max_row, column=1).font = Font(bold=True)
                texts = [block["text"]] if block["type"] == "paragraph" else block["items"]
                for text in texts:
                    summary.append([str(text)])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def render_pptx(model: Dict[str, Any]) -> bytes:
    """One slide per section; long tables continue on extra slides."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    navy = RGBColor(0x1F, 0x3A, 0x5F)

    def text_box(slide, left, top, width, height, text, size=14, bold=False, color=None):
        frame = slide.shapes.add_textbox(left, top, width, height).text_frame
        frame.word_wrap = True
        para = frame.paragraphs[0]
        para.text = str(text)
        para.font.size, para.font.bold = Pt(size), bold
        if color is not None:
            para.font.color.rgb = color
        return frame

    def new_slide(title):
        slide = prs.slides.add_slide(blank)
        text_box(slide, Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.8), title, size=26, bold=True, color=navy)
        return slide

    title = prs.slides.add_slide(blank)
    text_box(title, Inches(0.8), Inches(2.4), Inches(11.7), Inches(1.2), model["title"], size=40, bold=True, color=navy)
    text_box(title, Inches(0.8), Inches(3.6), Inches(11.7), Inches(0.6),
             f"{model['subtitle']} | Prepared for: {model['audience']}", size=18)
    if model["review_status"] != "ready_for_review" or model["warnings"]:
        text_box(title, Inches(0.8), Inches(5.6), Inches(11.7), Inches(1.2),
                 "Review required before distribution. " + ("; ".join(model["warnings"][:3]) or "Verify against sources."),
                 size=12, color=RGBColor(0x9C, 0x00, 0x06))

    rows_per_slide = 8
    for sec in model["sections"]:
        slide = new_slide(sec["title"])
        top = Inches(1.3)
        for block in sec["blocks"]:
            if block["type"] in ("paragraph", "bullets"):
                lines = ([f"{block['label']}: {block['text']}" if block.get("label") else block["text"]]
                         if block["type"] == "paragraph" else [f"• {item}" for item in block["items"]])
                frame = text_box(slide, Inches(0.6), top, Inches(12.1), Inches(0.5), lines[0], size=14)
                for line in lines[1:8]:
                    para = frame.add_paragraph()
                    para.text, para.font.size = line, Pt(14)
                top += Inches(0.45 * min(len(lines), 8) + 0.2)
            elif block["type"] == "table":
                chunks = [block["rows"][i:i + rows_per_slide] for i in range(0, len(block["rows"]), rows_per_slide)] or [[]]
                for index, chunk in enumerate(chunks):
                    if index:
                        slide, top = new_slide(f"{sec['title']} (continued)"), Inches(1.3)
                    shape = slide.shapes.add_table(len(chunk) + 1, len(block["columns"]), Inches(0.5), top,
                                                   Inches(12.3), Inches(0.4 * (len(chunk) + 1)))
                    table = shape.table
                    for col, name in enumerate(block["columns"]):
                        table.cell(0, col).text = str(name)
                    for r, row in enumerate(chunk, start=1):
                        for col, value in enumerate(row):
                            table.cell(r, col).text = str(value)[:180]
                    for r in range(len(chunk) + 1):
                        for col in range(len(block["columns"])):
                            for para in table.cell(r, col).text_frame.paragraphs:
                                para.font.size = Pt(11)
                    top += Inches(0.4 * (len(chunk) + 1) + 0.3)
            if top > Inches(6.6):
                slide, top = new_slide(f"{sec['title']} (continued)"), Inches(1.3)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


RENDERERS = {"docx": render_docx, "pdf": render_pdf, "xlsx": render_xlsx, "pptx": render_pptx}


async def render_formats(model: Dict[str, Any], formats: List[str]) -> Tuple[Dict[str, bytes], List[Dict[str, str]]]:
    """Render requested adaptive formats concurrently and validate each artifact."""
    from starlette.concurrency import run_in_threadpool

    from app.core.reports.report_guardrails import ReportGuardrails

    targets = [f for f in formats if f in RENDERERS]
    outputs = await asyncio.gather(*(run_in_threadpool(RENDERERS[f], model) for f in targets), return_exceptions=True)
    artifacts, errors = {}, []
    for fmt, out in zip(targets, outputs):
        if isinstance(out, Exception):
            errors.append({"format": fmt, "error": f"Generation failed ({type(out).__name__})."})
            continue
        if not ReportGuardrails.validate_artifact(fmt, out).get("valid"):
            errors.append({"format": fmt, "error": "Generated artifact failed structural validation."})
            continue
        artifacts[fmt] = out
    return artifacts, errors


# ── Orchestration ──────────────────────────────────────────────────────

async def build_plan(
    req: DocumentRequest, deal: Dict[str, Any], agent_results: List[Dict[str, Any]], registry: Any = None,
) -> Tuple[DocumentPlan, Dict[str, Any], List[Dict[str, Any]]]:
    """Plan without generating (cheap): evidence, risk register, coverage, architect ordering."""
    from app.core.knowledge_graph.service import risk_register as graph_risk_register
    from app.core.reports.document_planner import latest_agent_outputs
    from app.core.reports.evidence_brief import build_evidence_brief

    latest = latest_agent_outputs(agent_results)
    evidence = build_evidence_brief(deal, latest)
    risks = collect_risk_register(latest, await graph_risk_register(deal.get("id"), limit=25))
    plan = plan_document(req, deal, latest, evidence, risks)
    if req.use_architect:
        plan = await refine_with_architect(plan, registry, deal)
    return plan, evidence, risks


async def run_document_workflow(
    req: DocumentRequest, deal: Dict[str, Any], agent_results: List[Dict[str, Any]], registry: Any = None,
) -> Dict[str, Any]:
    """Full adaptive run. Returns plan, model, artifacts and errors (publishing is the caller's job)."""
    from app.core.reports.document_planner import prepare_document_payload

    plan, _, _ = await build_plan(req, deal, agent_results, registry)
    gap_results: List[Dict[str, Any]] = []
    if req.fill_gaps and plan.gaps:
        gap_results = await fill_gaps(plan, deal, registry, max_agents=req.max_gap_agents)
        if gap_results:
            agent_results = list(agent_results) + gap_results
            # Re-plan with the new evidence; sections may now be covered.
            plan, _, _ = await build_plan(req, deal, agent_results, registry)

    results, analyst_data, evidence = await prepare_document_payload(registry, deal, agent_results)
    from app.core.knowledge_graph.service import risk_register as graph_risk_register

    risks = collect_risk_register(results, await graph_risk_register(deal.get("id"), limit=25))
    model = compose_document_model(plan, deal, analyst_data, evidence, results, risks)
    analyst_data["risk_matrix"] = [
        {"risk": r["risk"], "severity": r["severity"], "category": r["category"],
         "mitigation": r.get("mitigation") or "", "evidence": ", ".join(r["sources"])}
        for r in risks[:20]
    ]
    artifacts, errors = await render_formats(model, plan.formats)
    review_actions = []
    if plan.gaps:
        review_actions.append("Resolve or accept the listed diligence gaps before distribution.")
    if model["review_status"] == "review_required":
        review_actions.append("Human review is required: verify material statements against cited sources.")
    if errors:
        review_actions.append("Some formats failed to render; regenerate or choose other formats.")
    return {
        "plan": plan.as_dict(),
        "model": model,
        "artifacts": artifacts,
        "errors": errors,
        "gap_fill": [{"agent": g["agent_type"], "task_id": g["task_id"]} for g in gap_results],
        "analyst_data": analyst_data,
        "agent_results": results,
        "evidence": evidence,
        "review_actions": review_actions,
    }
